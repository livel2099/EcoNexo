import json,re
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
r=Path(__file__).parent
rows=json.loads((r/'guion.json').read_text(encoding='utf-8'))
titles=['Decisiones para el campo de hoy','Cada decisión tiene un costo','La información por lote','Cultivos y ganadería','Ventanas para las labores','Seguimiento del agua y del cultivo','Presupuesto de forraje','Historial de las evaluaciones','Valor para cada equipo','El primer paso con un lote','Invitación a una demostración']
cues={3:'Señalar nombre del lote, cultivo y balance de agua. Aclarar que son datos de demostración.',5:'Señalar el aviso de pulverización y la necesidad de verificar las condiciones en el campo.',6:'Señalar los días y la relación entre lluvia y demanda del cultivo. Los números son ilustrativos.',7:'Señalar las mediciones de forraje, el aprovechamiento y la demanda animal.',8:'Señalar fecha, presupuesto y datos pendientes. El historial corresponde al ejemplo demostrativo.',10:'Señalar la vista de resumen y proponer trabajar con un lote del interesado.'}
trans=['Estas decisiones aparecen todos los días en el campo.','Veamos cómo reunimos esa información alrededor del lote.','La plataforma acompaña tanto cultivos como ganadería.','Primero, una tarea concreta: organizar una aplicación.','Además de una labor puntual, importa seguir la evolución del cultivo.','En ganadería, la pregunta cambia: cuánto alimento tenemos disponible.','Después del cálculo, necesitamos poder volver a revisar lo que evaluamos.','Ese seguimiento tiene valor para quienes trabajan juntos.','La mejor forma de evaluarlo es empezar con una necesidad propia.','Para cerrar, les propongo una pregunta.','Gracias. Quedo disponible para conversar al finalizar.']
d=Document();sec=d.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69);sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.8)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=d.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.space_after=Pt(7)
d.styles['Normal'].paragraph_format.line_spacing=1.12
d.styles['Title'].font.size=Pt(25);d.styles['Heading 1'].font.size=Pt(16);d.styles['Heading 2'].font.size=Pt(11)
d.add_paragraph('Guion de Econexo para la jornada Agtech','Title')
d.add_paragraph('Charla comercial de 10 a 15 minutos')
d.add_paragraph('Este guion acompaña las 11 diapositivas de la presentación con imágenes. Está pensado para explicar qué resuelve Econexo en el trabajo cotidiano del agro y cerrar con una invitación a una demostración sobre un lote propio.')
d.add_paragraph('Cómo usar el guion','Heading 1')
d.add_paragraph('Ensayá a ritmo conversado y dejá pausas breves. Las indicaciones de pantalla y las transiciones son ayudas para el orador. Las capturas muestran la interfaz del proyecto con datos de demostración; no representan resultados de un cliente.')
d.add_paragraph('Duración y adaptación','Heading 1')
d.add_paragraph('El recorrido sugerido dura aproximadamente 13 minutos. Para una versión de 10 minutos, abreviá los ejemplos de las diapositivas 5 a 8. Para llegar a 15 minutos, reservá el cierre para una pregunta del público.')
d.add_paragraph('Mensaje central','Heading 1')
d.add_paragraph('Econexo reúne información por lote para que el productor y su asesor puedan revisar las condiciones, planificar tareas y dar seguimiento a sus decisiones.')
for i,row in enumerate(rows):
 if i%2==0:d.add_page_break()
 n=i+1;notes=row['notes'];time=notes.split('\n')[0].replace('TIEMPO SUGERIDO: ','');body=notes.split('\n\nFuentes del producto:')[0].split('\n\n',1)[1]
 body=body.split('\n\nDuración total sugerida:')[0]
 if n==8:body='Cada evaluación que guardamos deja una referencia para volver a revisar. En la pantalla vemos un ejemplo de presupuesto forrajero con su fecha, los datos pendientes y la procedencia declarada. Eso permite saber qué información se usó y qué todavía necesita una verificación. Imaginemos una semana de trabajo: consultamos el lote, identificamos qué merece atención, lo contrastamos en una recorrida y decidimos con el asesor. Después podemos volver al historial para seguir la evolución y actualizar la evidencia. El valor está en sostener ese seguimiento con una referencia compartida. Los valores de esta imagen son de demostración y no constituyen una recomendación de carga para un campo real.'
 d.add_paragraph(f'{n:02d} {titles[i]}','Heading 1');d.add_paragraph('Tiempo sugerido '+time)
 if n in cues:
  p=d.add_paragraph();p.add_run('En pantalla  ').bold=True;p.add_run(cues[n])
 d.add_paragraph('Texto para decir','Heading 2');d.add_paragraph(body)
 p=d.add_paragraph();p.add_run('Transición  ').bold=True;p.add_run(trans[i])
f=sec.footer.paragraphs[0];f.alignment=2;run=f.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.addnext(field)
d.core_properties.title='Guion de Econexo para la jornada Agtech';d.core_properties.author='Econexo'
d.save(r.parent/'output/presentacion-agtech/Guion_Econexo_Jornada_Agtech.docx')
